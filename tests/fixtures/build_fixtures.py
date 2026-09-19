"""Builds `tests/fixtures/sample_essay.docx`: a short law essay with *real* Word
footnotes (a `word/footnotes.xml` part, a relationship, a content-type override,
and `w:footnoteReference` runs in the body).

python-docx has no footnote API, so the footnote part and the reference runs are
built by hand with lxml and spliced into a document python-docx produced, at the
raw zip level. The result must open cleanly in Word (and in python-docx).

Run directly to (re)generate the fixture::

    uv run python tests/fixtures/build_fixtures.py

Or import `build_sample_essay()` / `build_docx()` from tests.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import Any

import docx

FIXTURES_DIR = Path(__file__).parent
SAMPLE_ESSAY = FIXTURES_DIR / "sample_essay.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
NSMAP = {"w": W_NS}

TEST_AUTHOR = "Original Author"
TEST_DATE = "2020-01-01T00:00:00Z"

_id_counter = [9000]


def qn(tag: str) -> str:
    prefix, local = tag.split(":")
    assert prefix == "w"
    return f"{{{W_NS}}}{local}"


# --------------------------------------------------------------------------- #
# lxml element builders
# --------------------------------------------------------------------------- #

from lxml import etree  # noqa: E402


def _next_id() -> int:
    _id_counter[0] += 1
    return _id_counter[0]


def _rpr(italic: bool | str = False, style: str | None = None):
    rpr = etree.Element(qn("w:rPr"))
    if style:
        se = etree.SubElement(rpr, qn("w:rStyle"))
        se.set(qn("w:val"), style)
    if italic is True:
        etree.SubElement(rpr, qn("w:i"))
        etree.SubElement(rpr, qn("w:iCs"))
    if len(rpr) == 0:
        return None
    return rpr


def _run(text: str | None = None, italic: bool | str = False, style: str | None = None,
         tab: bool = False, br: bool = False):
    r = etree.Element(qn("w:r"))
    rpr = _rpr(italic, style)
    if rpr is not None:
        r.append(rpr)
    if tab:
        etree.SubElement(r, qn("w:tab"))
    elif br:
        etree.SubElement(r, qn("w:br"))
    elif text is not None:
        t = etree.SubElement(r, qn("w:t"))
        t.set(XML_SPACE, "preserve")
        t.text = text
    return r


def _content_runs(tokens: list[tuple]) -> list:
    """Build a list of <w:r>/<w:ins>/<w:del> elements from a token list.

    Token forms:
      ("text", str, italic)     italic is True / False / "emphasis"
      ("tab",)
      ("br",)
      ("del", str)              deleted text (w:delText inside w:del)
      ("ins", [tokens...])      inserted content (w:ins wrapping nested runs)
    """
    elems = []
    for tok in tokens:
        kind = tok[0]
        if kind == "text":
            _, text, italic = tok
            style = "Emphasis" if italic == "emphasis" else None
            ital = italic is True
            elems.append(_run(text=text, italic=ital, style=style))
        elif kind == "tab":
            elems.append(_run(tab=True))
        elif kind == "br":
            elems.append(_run(br=True))
        elif kind == "del":
            _, text = tok
            wrapper = etree.Element(qn("w:del"))
            wrapper.set(qn("w:id"), str(_next_id()))
            wrapper.set(qn("w:author"), TEST_AUTHOR)
            wrapper.set(qn("w:date"), TEST_DATE)
            r = etree.SubElement(wrapper, qn("w:r"))
            dt = etree.SubElement(r, qn("w:delText"))
            dt.set(XML_SPACE, "preserve")
            dt.text = text
            elems.append(wrapper)
        elif kind == "ins":
            _, nested = tok
            wrapper = etree.Element(qn("w:ins"))
            wrapper.set(qn("w:id"), str(_next_id()))
            wrapper.set(qn("w:author"), TEST_AUTHOR)
            wrapper.set(qn("w:date"), TEST_DATE)
            for sub in _content_runs(nested):
                wrapper.append(sub)
            elems.append(wrapper)
        else:  # pragma: no cover - fixture authoring error
            raise ValueError(f"unknown token kind: {kind!r}")
    return elems


def _body_footnote_reference_run(fn_id: int):
    """The clickable reference mark in the body: <w:r><w:footnoteReference/></w:r>."""
    r = etree.Element(qn("w:r"))
    rpr = etree.SubElement(r, qn("w:rPr"))
    se = etree.SubElement(rpr, qn("w:rStyle"))
    se.set(qn("w:val"), "FootnoteReference")
    ref = etree.SubElement(r, qn("w:footnoteReference"))
    ref.set(qn("w:id"), str(fn_id))
    return r


def _footnote_ref_marker():
    """The auto-numbered mark inside the footnote's own text: <w:footnoteRef/>."""
    r = etree.Element(qn("w:r"))
    rpr = etree.SubElement(r, qn("w:rPr"))
    se = etree.SubElement(rpr, qn("w:rStyle"))
    se.set(qn("w:val"), "FootnoteReference")
    etree.SubElement(r, qn("w:footnoteRef"))
    return r


def _footnote_element(fn_id: int, paragraphs: list[list[tuple]], pstyle: str = "FootnoteText"):
    fn = etree.Element(qn("w:footnote"))
    fn.set(qn("w:id"), str(fn_id))
    for i, tokens in enumerate(paragraphs):
        p = etree.SubElement(fn, qn("w:p"))
        ppr = etree.SubElement(p, qn("w:pPr"))
        pstyle_el = etree.SubElement(ppr, qn("w:pStyle"))
        pstyle_el.set(qn("w:val"), pstyle)
        if i == 0:
            p.append(_footnote_ref_marker())
            p.append(_run(text=" "))
        for el in _content_runs(tokens):
            p.append(el)
    return fn


def _separator_footnote(fn_id: int, kind: str):
    fn = etree.Element(qn("w:footnote"))
    fn.set(qn("w:id"), str(fn_id))
    fn.set(qn("w:type"), kind)
    p = etree.SubElement(fn, qn("w:p"))
    ppr = etree.SubElement(p, qn("w:pPr"))
    spacing = etree.SubElement(ppr, qn("w:spacing"))
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    r = etree.SubElement(p, qn("w:r"))
    marker_tag = "w:separator" if kind == "separator" else "w:continuationSeparator"
    etree.SubElement(r, qn(marker_tag))
    return fn


def _footnotes_xml_bytes(footnote_specs: list[tuple[int, list]]) -> bytes:
    root = etree.Element(qn("w:footnotes"), nsmap=NSMAP)
    root.append(_separator_footnote(-1, "separator"))
    root.append(_separator_footnote(0, "continuationSeparator"))
    for fn_id, paragraphs in footnote_specs:
        root.append(_footnote_element(fn_id, paragraphs))
    return b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + etree.tostring(root)


def _add_body_paragraph(document, tokens: list[tuple], style: str | None = None):
    p = document.add_paragraph(style=style)
    for tok in tokens:
        if tok[0] == "text":
            _, text, italic = tok
            run = p.add_run(text)
            if italic:
                run.italic = True
        elif tok[0] == "fnref":
            p._p.append(_body_footnote_reference_run(tok[1]))
        else:  # pragma: no cover
            raise ValueError(f"unknown body token: {tok!r}")
    return p


_EXTRA_STYLES = """
<w:style w:type="paragraph" w:styleId="FootnoteText">
<w:name w:val="footnote text"/><w:basedOn w:val="Normal"/><w:link w:val="FootnoteTextChar"/>
<w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>
<w:rPr><w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr>
</w:style>
<w:style w:type="character" w:styleId="FootnoteTextChar">
<w:name w:val="Footnote Text Char"/><w:basedOn w:val="DefaultParagraphFont"/><w:link w:val="FootnoteText"/>
</w:style>
<w:style w:type="character" w:styleId="FootnoteReference">
<w:name w:val="footnote reference"/><w:basedOn w:val="DefaultParagraphFont"/>
<w:rPr><w:vertAlign w:val="superscript"/></w:rPr>
</w:style>
""".strip()


# --------------------------------------------------------------------------- #
# Public builder
# --------------------------------------------------------------------------- #


def build_docx(dst: str | Path, body_paragraphs: list[tuple[str, Any]],
                footnote_specs: list[tuple[int, list]]) -> Path:
    """Build a .docx with real Word footnotes at `dst`.

    `body_paragraphs`: list of ("heading", text) | ("para", tokens) where tokens
    are the same ("text", str, italic) / ("fnref", id) tuples used elsewhere.
    `footnote_specs`: list of (footnote_id, paragraphs) as consumed by
    `_footnote_element` (paragraphs is a list of token-lists, one per w:p).
    """
    dst = Path(dst)
    document = docx.Document()
    for kind, payload in body_paragraphs:
        if kind == "heading":
            document.add_heading(payload, level=1)
        elif kind == "para":
            _add_body_paragraph(document, payload)
        else:  # pragma: no cover
            raise ValueError(f"unknown paragraph kind: {kind!r}")

    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    with zipfile.ZipFile(buf) as zin:
        names = zin.namelist()
        infos = {name: zin.getinfo(name) for name in names}
        data = {name: zin.read(name) for name in names}

    data["word/footnotes.xml"] = _footnotes_xml_bytes(footnote_specs)
    names.append("word/footnotes.xml")

    ct = data["[Content_Types].xml"].decode("utf-8")
    override = (
        '<Override PartName="/word/footnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
    )
    ct = ct.replace("</Types>", override + "</Types>")
    data["[Content_Types].xml"] = ct.encode("utf-8")

    rels_name = "word/_rels/document.xml.rels"
    rels = data[rels_name].decode("utf-8")
    existing_ids = [int(m) for m in re.findall(r'Id="rId(\d+)"', rels)]
    next_rid = max(existing_ids, default=0) + 1
    rel = (
        f'<Relationship Id="rId{next_rid}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
        'Target="footnotes.xml"/>'
    )
    rels = rels.replace("</Relationships>", rel + "</Relationships>")
    data[rels_name] = rels.encode("utf-8")

    styles_name = "word/styles.xml"
    styles = data[styles_name].decode("utf-8")
    styles = styles.replace("</w:styles>", _EXTRA_STYLES + "</w:styles>")
    data[styles_name] = styles.encode("utf-8")

    dst.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in names:
            info = infos.get(name) or zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            zout.writestr(info, data[name])
    return dst


# --------------------------------------------------------------------------- #
# The sample essay content
# --------------------------------------------------------------------------- #

FOOTNOTE_SPECS: list[tuple[int, list]] = [
    (1, [[  # case
        ("text", "Mabo v Queensland (No 2)", True),
        ("text", " (1992) 175 C.L.R. 1 at 42 per Brennan J.", False),
    ]]),
    (2, [[  # legislation
        ("text", "See ", False),
        ("text", "Competition and Consumer Act 2010 (Cth)", True),
        ("text", ", section 18.", False),
    ]]),
    (3, [[  # Ibid, with a pre-existing tracked change (40 -> 45)
        ("text", "Ibid, at ", False),
        ("del", "40"),
        ("ins", [("text", "45", False)]),
        ("text", ".", False),
    ]]),
    (4, [[  # above n
        ("text", "Mabo", True),
        ("text", ", above n 1, 60.", False),
    ]]),
    (5, [[  # two sources separated by ';'
        ("text", "Smith v Jones", True),
        ("text", " (2005) 220 A.L.R. 1; ", False),
        ("text", "Attorney-General (NSW) v X", True),
        ("text", " (2001) 53 N.S.W.L.R. 1.", False),
    ]]),
    (6, [  # journal article, split across two footnote paragraphs
        [("text", "J Smith, 'Native Title and the Common Law'", False)],
        [("text", "(1995) 19 ", False), ("text", "Melbourne University Law Review", True), ("text", " 195.", False)],
    ]),
    (7, [[  # book
        ("text", "P Butt, ", False),
        ("text", "Land Law", True),
        ("text", " (Lawbook Co, 6th ed, 2010) 45.", False),
    ]]),
    (8, [[  # commentary followed by a citation
        ("text", "This principle was later confirmed by the Court: ", False),
        ("text", "Wik Peoples v Queensland", True),
        ("text", " (1996) 187 C.L.R. 1, 129 per Toohey J.", False),
    ]]),
    (9, [[  # website with a URL, and a tab
        ("text", "Australian Human Rights Commission,", False),
        ("tab",),
        ("text", " 'Native Title Report 2017' (Web Page, 2017) "
                  "<https://humanrights.gov.au/native-title-report-2017>.", False),
    ]]),
    (10, [[  # report, italic via the "Emphasis" character style, and a line break
        ("text", "Australian Law Reform Commission, ", False),
        ("text", "Connection to Country: Review of the Native Title Act 1993 (Cth)", "emphasis"),
        ("text", " (Report No 126, 2015)", False),
        ("br",),
        ("text", "34.", False),
    ]]),
]

BODY_PARAGRAPHS: list[tuple[str, Any]] = [
    ("heading", "Native Title and the Common Law of Australia"),
    ("para", [
        ("text", "The recognition of native title in Australian law marks one of the most significant "
                  "developments in the modern common law of Australia.", False),
        ("fnref", 1),
        ("text", " Prior to this recognition, colonial courts proceeded on the assumption that the "
                  "continent was legally unoccupied at the time of settlement, a position increasingly "
                  "difficult to reconcile with statutes such as the Commonwealth's consumer protection "
                  "framework.", False),
        ("fnref", 2),
    ]),
    ("para", [
        ("text", "The High Court's reasoning has since been applied and refined in later proceedings.", False),
        ("fnref", 3),
        ("text", " Commentators continue to refer to the decision, for short, simply as Mabo.", False),
        ("fnref", 4),
        ("text", " Its influence can also be traced through subsequent litigation concerning "
                  "overlapping claims.", False),
        ("fnref", 5),
    ]),
    ("para", [
        ("text", "Academic commentary followed swiftly.", False),
        ("fnref", 6),
        ("text", " Textbook treatments were revised accordingly.", False),
        ("fnref", 7),
        ("text", " and the doctrine was tested again a few years later.", False),
        ("fnref", 8),
    ]),
    ("para", [
        ("text", "Government and civil-society bodies have since monitored the practical operation of "
                  "native title.", False),
        ("fnref", 9),
        ("text", " including through periodic statutory review.", False),
        ("fnref", 10),
    ]),
]

# Expected result of read_footnotes(SAMPLE_ESSAY), for tests: (full_text, [(run_text, italic), ...])
EXPECTED_FOOTNOTES: list[tuple[str, list[tuple[str, bool]]]] = [
    ("Mabo v Queensland (No 2) (1992) 175 C.L.R. 1 at 42 per Brennan J.", [
        ("Mabo v Queensland (No 2)", True),
        (" (1992) 175 C.L.R. 1 at 42 per Brennan J.", False),
    ]),
    ("See Competition and Consumer Act 2010 (Cth), section 18.", [
        ("See ", False),
        ("Competition and Consumer Act 2010 (Cth)", True),
        (", section 18.", False),
    ]),
    ("Ibid, at 45.", [
        ("Ibid, at 45.", False),
    ]),
    ("Mabo, above n 1, 60.", [
        ("Mabo", True),
        (", above n 1, 60.", False),
    ]),
    ("Smith v Jones (2005) 220 A.L.R. 1; Attorney-General (NSW) v X (2001) 53 N.S.W.L.R. 1.", [
        ("Smith v Jones", True),
        (" (2005) 220 A.L.R. 1; ", False),
        ("Attorney-General (NSW) v X", True),
        (" (2001) 53 N.S.W.L.R. 1.", False),
    ]),
    ("J Smith, 'Native Title and the Common Law'\n(1995) 19 Melbourne University Law Review 195.", [
        ("J Smith, 'Native Title and the Common Law'\n(1995) 19 ", False),
        ("Melbourne University Law Review", True),
        (" 195.", False),
    ]),
    ("P Butt, Land Law (Lawbook Co, 6th ed, 2010) 45.", [
        ("P Butt, ", False),
        ("Land Law", True),
        (" (Lawbook Co, 6th ed, 2010) 45.", False),
    ]),
    ("This principle was later confirmed by the Court: Wik Peoples v Queensland (1996) 187 C.L.R. 1, "
     "129 per Toohey J.", [
        ("This principle was later confirmed by the Court: ", False),
        ("Wik Peoples v Queensland", True),
        (" (1996) 187 C.L.R. 1, 129 per Toohey J.", False),
    ]),
    ("Australian Human Rights Commission,\t 'Native Title Report 2017' (Web Page, 2017) "
     "<https://humanrights.gov.au/native-title-report-2017>.", [
        ("Australian Human Rights Commission,\t 'Native Title Report 2017' (Web Page, 2017) "
         "<https://humanrights.gov.au/native-title-report-2017>.", False),
    ]),
    ("Australian Law Reform Commission, Connection to Country: Review of the Native Title Act 1993 (Cth) "
     "(Report No 126, 2015)\n34.", [
        ("Australian Law Reform Commission, ", False),
        ("Connection to Country: Review of the Native Title Act 1993 (Cth)", True),
        (" (Report No 126, 2015)\n34.", False),
    ]),
]


def build_sample_essay(dst: str | Path | None = None) -> Path:
    return build_docx(dst or SAMPLE_ESSAY, BODY_PARAGRAPHS, FOOTNOTE_SPECS)


if __name__ == "__main__":
    path = build_sample_essay()
    print(f"wrote {path}")
