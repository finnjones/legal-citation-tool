"""Tests for aglc.docx_io: reading and writing real Word footnotes."""

from __future__ import annotations

import zipfile

import docx
import pytest
from lxml import etree

from aglc.docx_io import (
    DOCUMENT_PATH,
    FOOTNOTES_PATH,
    _footnote_elements,
    _parse_xml,
    qn,
    read_footnotes,
    write_document,
)
from aglc.models import BibliographySection, FootnoteResult, RichText

from tests.fixtures.build_fixtures import EXPECTED_FOOTNOTES, build_sample_essay


@pytest.fixture(scope="session")
def sample_essay(tmp_path_factory) -> str:
    dst = tmp_path_factory.mktemp("fixtures") / "sample_essay.docx"
    build_sample_essay(dst)
    return str(dst)


def _rt(*parts: tuple[str, bool]) -> RichText:
    r = RichText()
    for text, italic in parts:
        r.append(text, italic)
    return r


# --------------------------------------------------------------------------- #
# read_footnotes
# --------------------------------------------------------------------------- #


def test_fixture_opens_in_python_docx(sample_essay):
    # A minimal sanity check that the hand-built fixture is well-formed OOXML.
    d = docx.Document(sample_essay)
    assert len(d.paragraphs) == 5


def test_read_footnotes_count_and_numbering(sample_essay):
    footnotes = read_footnotes(sample_essay)
    assert len(footnotes) == 10
    assert [fn.number for fn in footnotes] == list(range(1, 11))


def test_read_footnotes_skips_separator_and_continuation(sample_essay):
    # If separator (-1) / continuationSeparator (0) leaked in, we'd have 12, not 10.
    footnotes = read_footnotes(sample_essay)
    assert len(footnotes) == len(EXPECTED_FOOTNOTES)


@pytest.mark.parametrize("index", range(10))
def test_read_footnotes_text_and_italics(sample_essay, index):
    footnotes = read_footnotes(sample_essay)
    exp_text, exp_runs = EXPECTED_FOOTNOTES[index]
    fn = footnotes[index]
    assert fn.original.text == exp_text
    assert [(r.text, r.italic) for r in fn.original.runs] == exp_runs


def test_read_footnotes_strips_single_leading_space_only(sample_essay):
    # fn1's content run starts with "M", not a lingering space; the auto space
    # Word inserts after the reference mark must be fully consumed.
    footnotes = read_footnotes(sample_essay)
    assert footnotes[0].original.text.startswith("Mabo")


def test_read_footnotes_multi_paragraph_joined_with_newline(sample_essay):
    footnotes = read_footnotes(sample_essay)
    fn6 = footnotes[5]
    assert "\n" in fn6.original.text
    assert fn6.original.text.startswith("J Smith, 'Native Title and the Common Law'\n")


def test_read_footnotes_tab_and_break_characters(sample_essay):
    footnotes = read_footnotes(sample_essay)
    assert "\t" in footnotes[8].original.text  # website footnote has a w:tab
    assert "\n" in footnotes[9].original.text  # report footnote has a w:br


def test_read_footnotes_ignores_deleted_includes_inserted(sample_essay):
    footnotes = read_footnotes(sample_essay)
    fn3 = footnotes[2]
    assert fn3.original.text == "Ibid, at 45."
    assert "40" not in fn3.original.text


def test_read_footnotes_emphasis_style_counts_as_italic(sample_essay):
    footnotes = read_footnotes(sample_essay)
    fn10 = footnotes[9]
    italic_runs = [r for r in fn10.original.runs if r.italic]
    assert any("Connection to Country" in r.text for r in italic_runs)


# --------------------------------------------------------------------------- #
# write_document
# --------------------------------------------------------------------------- #


def _sample_results(sample_essay) -> list[FootnoteResult]:
    footnotes = read_footnotes(sample_essay)
    formatted_overrides = {
        1: _rt(("Mabo v Queensland [No 2]", True), (" (1992) 175 CLR 1, 42 (Brennan J).", False)),
        5: _rt(
            ("Smith v Jones", True),
            (" (2005) 220 ALR 1; ", False),
            ("Attorney-General (NSW) v X", True),
            (" (2001) 53 NSWLR 1.", False),
        ),
        6: _rt(
            ("Smith, 'Native Title and the Common Law'", False),
            (" (1995) 19 ", False),
            ("Melbourne University Law Review", True),
            (" 195.", False),
        ),
        10: _rt(
            ("Australian Law Reform Commission, ", False),
            ("Connection to Country", True),
            (" (Report No 126, 2015) 34.", False),
        ),
    }
    results = []
    for fn in footnotes:
        formatted = formatted_overrides.get(fn.number, fn.original)
        results.append(FootnoteResult(number=fn.number, original=fn.original, formatted=formatted))
    return results


def _sample_bibliography() -> list[BibliographySection]:
    return [
        BibliographySection(
            heading="A Articles/Books/Reports",
            entries=[_rt(("Butt, P, ", False), ("Land Law", True), (" (Lawbook Co, 6th ed, 2010)", False))],
        ),
        BibliographySection(
            heading="B Cases",
            entries=[_rt(("Mabo v Queensland [No 2]", True), (" (1992) 175 CLR 1", False))],
        ),
    ]


@pytest.mark.parametrize("track_changes", [True, False])
def test_write_document_round_trip(tmp_path, sample_essay, track_changes):
    results = _sample_results(sample_essay)
    bibliography = _sample_bibliography()
    dst = tmp_path / f"out-{track_changes}.docx"

    write_document(
        sample_essay, dst, results, bibliography=bibliography, track_changes=track_changes, author="AGLC Bot"
    )

    # Valid zip a la python-docx.
    d = docx.Document(str(dst))
    assert d is not None

    # Re-reading recovers the *accepted* text of every footnote (delText ignored,
    # ins included, deleted paragraph marks don't leave a stray "\n").
    reread = read_footnotes(str(dst))
    by_number = {fn.number: fn for fn in reread}
    for result in results:
        assert by_number[result.number].original.text == result.formatted.text


def test_write_document_tracked_changes_have_author_and_date(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    dst = tmp_path / "tracked.docx"
    write_document(sample_essay, dst, results, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(dst) as z:
        root = _parse_xml(z.read(FOOTNOTES_PATH))

    ins_els = list(root.iter(qn("w:ins")))
    del_els = list(root.iter(qn("w:del")))
    assert ins_els and del_els

    ids_seen = set()
    for el in ins_els + del_els:
        assert el.get(qn("w:author")) is not None
        date = el.get(qn("w:date"))
        assert date is not None and date.endswith("Z")  # ISO 8601 UTC
        fid = el.get(qn("w:id"))
        assert fid is not None
        assert fid not in ids_seen  # every w:ins/w:del id is unique
        ids_seen.add(fid)


def test_write_document_new_changes_have_our_author(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    dst = tmp_path / "tracked.docx"
    write_document(sample_essay, dst, results, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(dst) as z:
        root = _parse_xml(z.read(FOOTNOTES_PATH))
    authors = {el.get(qn("w:author")) for el in root.iter(qn("w:ins"))}
    authors |= {el.get(qn("w:author")) for el in root.iter(qn("w:del"))}
    assert "AGLC Bot" in authors


def test_write_document_clean_mode_has_no_new_tracked_changes(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    dst = tmp_path / "clean.docx"
    write_document(sample_essay, dst, results, track_changes=False, author="AGLC Bot")

    with zipfile.ZipFile(dst) as z:
        root = _parse_xml(z.read(FOOTNOTES_PATH))
    fn_map = _footnote_elements(root)

    # The footnotes we actually rewrote must contain no w:ins/w:del at all.
    for number in (1, 5, 6, 10):
        xml = etree.tostring(fn_map[str(number)])
        assert b"<w:ins" not in xml
        assert b"<w:del" not in xml

    with zipfile.ZipFile(dst) as z:
        doc_xml = z.read(DOCUMENT_PATH)
    assert b"<w:ins" not in doc_xml
    assert b"<w:del" not in doc_xml


def test_write_document_unchanged_footnotes_are_byte_identical(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    dst = tmp_path / "tracked.docx"
    write_document(sample_essay, dst, results, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(sample_essay) as z:
        src_root = _parse_xml(z.read(FOOTNOTES_PATH))
    with zipfile.ZipFile(dst) as z:
        out_root = _parse_xml(z.read(FOOTNOTES_PATH))

    src_map = _footnote_elements(src_root)
    out_map = _footnote_elements(out_root)

    changed_numbers = {str(r.number) for r in results if r.changed}
    for fid, src_elem in src_map.items():
        if fid in changed_numbers:
            continue
        assert etree.tostring(src_elem) == etree.tostring(out_map[fid]), f"footnote {fid} should be untouched"


def test_write_document_changed_footnotes_differ(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    dst = tmp_path / "tracked.docx"
    write_document(sample_essay, dst, results, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(sample_essay) as z:
        src_root = _parse_xml(z.read(FOOTNOTES_PATH))
    with zipfile.ZipFile(dst) as z:
        out_root = _parse_xml(z.read(FOOTNOTES_PATH))
    src_map = _footnote_elements(src_root)
    out_map = _footnote_elements(out_root)

    for r in results:
        if r.changed:
            fid = str(r.number)
            assert etree.tostring(src_map[fid]) != etree.tostring(out_map[fid])


def test_write_document_bibliography_appears(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    bibliography = _sample_bibliography()
    dst = tmp_path / "tracked.docx"
    write_document(sample_essay, dst, results, bibliography=bibliography, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(dst) as z:
        doc_xml = z.read(DOCUMENT_PATH)
    root = _parse_xml(doc_xml)
    texts = [t.text for t in root.iter(qn("w:t"))]
    joined = " ".join(t for t in texts if t)

    assert "Bibliography" in joined
    assert "A Articles/Books/Reports" in joined
    assert "B Cases" in joined
    assert "Land Law" in joined
    assert 'w:type="page"' in etree.tostring(root).decode()


def test_write_document_bibliography_italics_preserved(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    bibliography = _sample_bibliography()
    dst = tmp_path / "tracked.docx"
    write_document(sample_essay, dst, results, bibliography=bibliography, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(dst) as z:
        doc_xml = z.read(DOCUMENT_PATH)
    root = _parse_xml(doc_xml)

    italic_texts = set()
    for r in root.iter(qn("w:r")):
        t = r.find(qn("w:t"))
        if t is None or not t.text:
            continue
        rpr = r.find(qn("w:rPr"))
        if rpr is not None and rpr.find(qn("w:i")) is not None:
            italic_texts.add(t.text)
    assert "Land Law" in italic_texts


def test_write_document_no_bibliography_when_not_given(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    dst = tmp_path / "no-bib.docx"
    write_document(sample_essay, dst, results, bibliography=None, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(dst) as z:
        doc_xml = z.read(DOCUMENT_PATH)
    assert b"Bibliography" not in doc_xml


def test_write_document_preserves_other_parts_untouched(tmp_path, sample_essay):
    results = _sample_results(sample_essay)
    dst = tmp_path / "tracked.docx"
    write_document(sample_essay, dst, results, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(sample_essay) as z:
        src_names = set(z.namelist())
        src_styles = z.read("word/styles.xml")
    with zipfile.ZipFile(dst) as z:
        out_names = set(z.namelist())
        out_styles = z.read("word/styles.xml")

    assert src_names == out_names
    assert src_styles == out_styles


def test_write_document_no_changes_is_a_noop_for_footnotes(tmp_path, sample_essay):
    footnotes = read_footnotes(sample_essay)
    results = [FootnoteResult(number=fn.number, original=fn.original, formatted=fn.original) for fn in footnotes]
    dst = tmp_path / "noop.docx"
    write_document(sample_essay, dst, results, track_changes=True, author="AGLC Bot")

    with zipfile.ZipFile(sample_essay) as z:
        src_fn = z.read(FOOTNOTES_PATH)
    with zipfile.ZipFile(dst) as z:
        out_fn = z.read(FOOTNOTES_PATH)
    assert src_fn == out_fn
